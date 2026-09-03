# -*- coding: utf-8 -*-
"""将 interview-cheatsheet.md 转为 Word 文档（保持中文排版、加粗、行内代码、引用、列表、表格）。"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

SRC = r"e:\github项目\ai-ink\interview-cheatsheet.md"
DST = r"e:\github项目\ai-ink\interview-cheatsheet.docx"
# 支持命令行参数：python md2docx.py [源.md] [目标.docx]

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_font(style, name_latin="Calibri", name_cjk="微软雅黑"):
    style.font.name = name_latin
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name_cjk)


def add_runs(p, text):
    """把一段 markdown 文本解析成 bold + code + 普通 run。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = p.add_run(part)
    return p


def add_rule(doc):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom)
    ppr.append(pbdr)


def parse_table(rows):
    def split(row):
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        return cells

    header = split(rows[0])
    data = [split(r) for r in rows[2:] if r.strip()]  # 跳过 |---| 分隔行
    return header, data


def main():
    global SRC, DST
    if len(sys.argv) >= 3:
        SRC = sys.argv[1]
        DST = sys.argv[2]
    elif len(sys.argv) == 2:
        SRC = sys.argv[1]
        DST = sys.argv[1][:-3] + ".docx"
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 全局样式
    normal = doc.styles["Normal"]
    set_font(normal)
    normal.font.size = Pt(10.5)
    for h, sz in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        set_font(doc.styles[h])
        doc.styles[h].font.size = Pt(sz)
        doc.styles[h].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    in_table = False
    table_rows = []

    def flush_table():
        if not table_rows:
            return
        header, data = parse_table(table_rows)
        t = doc.add_table(rows=1 + len(data), cols=len(header))
        t.style = "Table Grid"
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            add_runs(cell.paragraphs[0], h)
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, row in enumerate(data):
            for j, v in enumerate(row):
                if j < len(header):
                    cell = t.rows[i + 1].cells[j]
                    cell.text = ""
                    add_runs(cell.paragraphs[0], v)

    for line in lines:
        # 表格段
        if line.lstrip().startswith("|") and not in_table:
            in_table = True
            table_rows = [line]
            continue
        if in_table:
            if line.lstrip().startswith("|"):
                table_rows.append(line)
                continue
            else:
                flush_table()
                table_rows = []
                in_table = False

        if not line.strip():
            continue

        stripped = line.strip()

        # 标题
        if stripped.startswith("### "):
            p = doc.add_heading(level=3)
            add_runs(p, stripped[4:])
        elif stripped.startswith("## "):
            p = doc.add_heading(level=2)
            add_runs(p, stripped[3:])
        elif stripped.startswith("# "):
            p = doc.add_heading(level=1)
            add_runs(p, stripped[2:])
        elif stripped == "---":
            add_rule(doc)
        # 引用块
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.3)
            ppr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "4")
            left.set(qn("w:color"), "1F497D")
            pbdr.append(left)
            ppr.append(pbdr)
            add_runs(p, stripped[1:].strip())
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        # 列表（含嵌套缩进）
        elif re.match(r"^(\s*)[-*] ", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.25 + 0.25 * (indent // 2))
            add_runs(p, "•  " + line.lstrip()[2:])
        elif re.match(r"^(\s*)\d+\. ", line):
            indent = len(line) - len(line.lstrip())
            m = re.match(r"^(\s*)(\d+)\. (.*)$", line)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.25 + 0.25 * (indent // 2))
            add_runs(p, f"{m.group(2)}.  {m.group(3)}")
        # 普通段落
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)

    flush_table()

    doc.save(DST)
    print("saved:", DST)


if __name__ == "__main__":
    main()
